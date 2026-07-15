"""Word 成绩分析报告的构建与导出逻辑。"""

from __future__ import annotations

from datetime import date
from io import BytesIO
import logging
from pathlib import Path
import re
from tempfile import TemporaryDirectory
from zipfile import ZipFile

import pandas as pd
from docx.shared import Mm
from docxtpl import DocxTemplate, InlineImage

from grade_logic import is_total_score_column


LOGGER = logging.getLogger(__name__)
DEFAULT_TEMPLATE_PATH = Path(__file__).resolve().parent / "templates" / "score_report_template.docx"
_ILLEGAL_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]+')


class ReportGenerationError(RuntimeError):
    """面向页面展示的报告生成错误。"""


def _format_score(value: object) -> str:
    """按成绩展示习惯格式化数值，避免无意义的小数点。"""
    number = float(value)
    return str(int(number)) if number.is_integer() else f"{number:.2f}".rstrip("0").rstrip(".")


def _format_rate(value: object) -> str:
    return f"{float(value):.1f}%"


def _format_student_list(students: pd.DataFrame, empty_message: str) -> str:
    if students is None or students.empty:
        return empty_message

    return "\n".join(
        f"{row['姓名']}：{_format_score(row['分数'])} 分"
        for _, row in students.iterrows()
    )


def _build_summary_text(
    analysis_result: dict, selected_class: str, score_col: str, full_score: float
) -> str:
    student_count = int(analysis_result["student_count"])
    average_score = float(analysis_result["average_score"])
    highest_score = float(analysis_result["highest_score"])
    lowest_score = float(analysis_result["lowest_score"])
    average_rate = average_score / float(full_score) * 100 if float(full_score) > 0 else 0.0
    score_range = highest_score - lowest_score

    if average_rate >= 85:
        level_description = "整体表现较好"
    elif average_rate >= 70:
        level_description = "整体表现中等偏上"
    elif average_rate >= 60:
        level_description = "整体基本达到要求"
    elif average_rate >= 40:
        level_description = "整体仍有较大提升空间"
    else:
        level_description = "基础层学生占比较大，应优先夯实基础"

    paragraphs = [
        (
            f"本次{selected_class}的{score_col}分析共覆盖 {student_count} 名参考学生。"
            f"平均分为 {_format_score(average_score)} 分，当前分析列满分为 "
            f"{_format_score(full_score)} 分，平均得分率为 {average_rate:.1f}%。"
        ),
        (
            f"最高分为 {_format_score(highest_score)} 分，最低分为 "
            f"{_format_score(lowest_score)} 分，极差为 {_format_score(score_range)} 分；"
            f"及格率为 {_format_rate(analysis_result['pass_rate'])}，优秀率为 "
            f"{_format_rate(analysis_result['excellent_rate'])}。"
        ),
        f"按平均得分率进行保守判断，当前{level_description}。具体结构需结合后续分布与等级图共同观察。",
    ]
    if student_count < 10:
        paragraphs.append("本次样本较少，结论仅供参考，不宜据此作过度推断。")
    if is_total_score_column(score_col):
        paragraphs.append("总分分析反映整体结果，具体薄弱学科仍需结合单科数据进一步判断。")
    return "\n".join(paragraphs)


def _valid_distribution_rows(distribution: pd.DataFrame) -> pd.DataFrame:
    """保留网页分布结果中可直接解释的非空区间，不重新计算成绩。"""
    required_columns = {"档位", "区间", "人数", "占比"}
    if distribution is None or distribution.empty or not required_columns.issubset(distribution.columns):
        return pd.DataFrame(columns=["档位", "区间", "人数", "占比"])

    rows = distribution.loc[:, ["档位", "区间", "人数", "占比"]].copy()
    rows["人数"] = pd.to_numeric(rows["人数"], errors="coerce")
    rows["占比"] = pd.to_numeric(rows["占比"], errors="coerce")
    rows["档位"] = rows["档位"].astype(str).str.strip()
    rows["区间"] = rows["区间"].astype(str).str.strip()
    return rows[
        rows["人数"].notna()
        & rows["占比"].notna()
        & (rows["人数"] > 0)
        & (rows["档位"] != "")
        & (rows["区间"] != "")
        & (rows["区间"] != "无区间")
    ].reset_index(drop=True)


def _build_distribution_analysis(distribution: pd.DataFrame) -> str:
    rows = _valid_distribution_rows(distribution)
    if rows.empty:
        return "当前没有可用于分布判断的有效区间数据。"

    maximum_count = rows["人数"].max()
    dominant_rows = rows[rows["人数"] == maximum_count]
    dominant_intervals = dominant_rows["区间"].tolist()
    if len(dominant_intervals) == 1:
        interval = dominant_intervals[0]
        share = float(dominant_rows.iloc[0]["占比"])
        if share >= 50:
            dominant_text = (
                f"主体分布：成绩明显集中于{interval}区间，共 {int(maximum_count)} 人，"
                f"占参考人数的 {share:.1f}%。"
            )
        else:
            dominant_text = (
                f"主体分布：人数最多的分数区间为{interval}，共 {int(maximum_count)} 人，"
                f"占参考人数的 {share:.1f}%。"
            )
    else:
        interval_names = "、".join(dominant_intervals)
        shares = dominant_rows["占比"].astype(float).tolist()
        share_text = f"均占 {shares[0]:.1f}%" if len(set(round(x, 6) for x in shares)) == 1 else "占比相近"
        interval_count_text = {2: "两个", 3: "三个"}.get(
            len(dominant_intervals), f"{len(dominant_intervals)}个"
        )
        dominant_text = (
            f"主体分布：成绩主要集中在{interval_names}{interval_count_text}区间，"
            f"各 {int(maximum_count)} 人，{share_text}。"
        )

    first_row = rows.iloc[0]
    last_row = rows.iloc[-1]
    edge_parts = [
        f"最低分数区间{first_row['区间']}有 {int(first_row['人数'])} 人，占 {float(first_row['占比']):.1f}%"
    ]
    if last_row.name != first_row.name:
        edge_parts.append(
            f"最高分数区间{last_row['区间']}有 {int(last_row['人数'])} 人，占 {float(last_row['占比']):.1f}%"
        )
    edge_text = "两端情况：" + "；".join(edge_parts) + "。"

    dominant_share = float(dominant_rows["占比"].max())
    low_share = float(first_row["占比"])
    high_share = float(last_row["占比"])
    middle_shares = rows.iloc[1:-1]["占比"].astype(float).tolist()
    is_polarized = (
        dominant_share < 50
        and len(rows) >= 3
        and low_share >= 25
        and high_share >= 25
        and low_share + high_share >= 60
        and (not middle_shares or max(middle_shares) < min(low_share, high_share))
    )
    if dominant_share >= 50:
        structure_text = "结构判断：单一区间占比达到一半以上，成绩分布呈明显集中。"
    elif is_polarized:
        structure_text = (
            "结构判断：最低与最高分数区间均占较高比例，两端合计达到六成以上，"
            "且各中间区间占比更低，呈现两极分化特征。"
        )
    else:
        structure_text = "结构判断：除主体区间外，其他区间仍有一定人数，整体分布较分散。"

    if low_share >= 20:
        focus_text = "关注方向：建议优先跟踪低分数区间学生的阶段变化，同时为主体区间学生设置可达成的进阶目标。"
    elif high_share >= 30:
        focus_text = "关注方向：建议保持高分数区间学生的稳定性，并持续推动中间区间学生向上迁移。"
    else:
        focus_text = "关注方向：建议围绕主体区间开展分层反馈，并关注相邻区间学生的短周期变化。"
    return "\n".join([dominant_text, edge_text, structure_text, focus_text])


def _level_rows(distribution: pd.DataFrame) -> pd.DataFrame:
    required_columns = {"档位", "人数", "占比"}
    if distribution is None or distribution.empty or not required_columns.issubset(distribution.columns):
        return pd.DataFrame(columns=["档位", "人数", "占比"])
    rows = distribution.loc[:, ["档位", "人数", "占比"]].copy()
    rows["人数"] = pd.to_numeric(rows["人数"], errors="coerce")
    rows["占比"] = pd.to_numeric(rows["占比"], errors="coerce")
    rows["档位"] = rows["档位"].astype(str).str.strip()
    return rows[rows["人数"].notna() & rows["占比"].notna() & (rows["档位"] != "")]


def _join_chinese_items(items: list[str]) -> str:
    if len(items) <= 1:
        return "".join(items)
    if len(items) == 2:
        return "和".join(items)
    return "、".join(items[:-1]) + "和" + items[-1]


def _build_level_analysis(distribution: pd.DataFrame) -> str:
    rows = _level_rows(distribution)
    if rows.empty:
        return "当前没有可用于等级结构判断的有效数据。"

    records = {row["档位"]: row for _, row in rows.iterrows()}
    ordered_levels = ["待提升", "及格", "中等", "良好", "优秀"]
    composition_parts = []
    for level in ordered_levels:
        row = records.get(level)
        if row is None:
            continue
        count = int(row["人数"])
        share = float(row["占比"])
        if count == 0:
            composition_parts.append(f"{level}层暂未形成（0人，{share:.1f}%）")
        else:
            composition_parts.append(f"{level}层 {count} 人（{share:.1f}%）")
    composition_text = "等级构成：" + "；".join(composition_parts) + "。"

    maximum_count = rows["人数"].max()
    dominant_levels = [
        level
        for level in ordered_levels
        if level in records and float(records[level]["人数"]) == maximum_count
    ]
    if len(dominant_levels) == 1:
        dominant_text = f"主体等级为{dominant_levels[0]}，是当前人数最多的层级。"
    else:
        level_count_text = {2: "两个", 3: "三个"}.get(
            len(dominant_levels), f"{len(dominant_levels)}个"
        )
        dominant_text = (
            f"主体等级为{_join_chinese_items(dominant_levels)}，"
            f"{level_count_text}层级人数并列最多。"
        )

    good_share = float(records.get("良好", {}).get("占比", 0.0))
    excellent_share = float(records.get("优秀", {}).get("占比", 0.0))
    middle_high_share = good_share + excellent_share
    if middle_high_share >= 50:
        formation_text = f"良好与优秀合计占 {middle_high_share:.1f}%，中高分层已形成一定规模。"
    elif middle_high_share >= 30:
        formation_text = f"良好与优秀合计占 {middle_high_share:.1f}%，中高分层已初步形成。"
    else:
        formation_text = f"良好与优秀合计占 {middle_high_share:.1f}%，中高分层规模仍有限。"

    improve_share = float(records.get("待提升", {}).get("占比", 0.0))
    pass_share = float(records.get("及格", {}).get("占比", 0.0))
    middle_share = float(records.get("中等", {}).get("占比", 0.0))
    if improve_share >= 20:
        priority_text = "下一阶段可优先降低待提升层占比，并帮助临界学生逐步进入及格及以上层级。"
    elif pass_share + middle_share >= 50:
        priority_text = "下一阶段可优先推动及格与中等层学生稳定向上迁移，同时保持中高分层的学习状态。"
    elif improve_share > 0:
        priority_text = "下一阶段可对少量待提升学生保持跟踪，并巩固现有中高分层结构。"
    else:
        priority_text = "下一阶段可在保持无待提升学生状态的同时，继续提升良好与优秀层的稳定性。"
    return "\n".join([composition_text, dominant_text + formation_text, priority_text])


def _build_teaching_suggestions(
    analysis_result: dict, distribution: pd.DataFrame, score_col: str
) -> str:
    records = {row["档位"]: row for _, row in _level_rows(distribution).iterrows()}
    improve_count = int(records.get("待提升", {}).get("人数", analysis_result.get("fail_count", 0)))
    pass_count = int(records.get("及格", {}).get("人数", 0))
    middle_high_count = sum(
        int(records.get(level, {}).get("人数", 0)) for level in ("中等", "良好", "优秀")
    )

    if improve_count:
        first = f"1. 对 {improve_count} 名待提升学生采用小目标、分层练习和阶段反馈，持续观察其变化。"
    else:
        first = "1. 当前未形成待提升层，建议继续保留基础任务检查，防止学习状态出现明显波动。"
    if pass_count:
        second = f"2. 对 {pass_count} 名及格层临界学生设置短周期、可检查的提升目标，重点巩固成绩稳定性。"
    else:
        second = "2. 结合相邻等级学生的变化设置短周期目标，避免只依据一次结果作长期判断。"
    if middle_high_count:
        third = f"3. 对 {middle_high_count} 名中等及以上学生兼顾巩固与适度拓展，保持不同层级的进阶空间。"
    else:
        third = "3. 中高分层尚未形成时，建议先稳定基础层与临界层，再逐步增加拓展性任务。"
    suggestions = [first, second, third]
    if is_total_score_column(score_col):
        suggestions.append("4. 后续教学判断应结合单科数据进一步判断，避免仅凭总分定位具体薄弱环节。")
    return "\n".join(suggestions)


def build_report_context(
    *,
    analysis_result: dict,
    excellent_df: pd.DataFrame,
    fail_df: pd.DataFrame,
    distribution: pd.DataFrame,
    selected_class: str,
    score_col: str,
    full_score: float,
    school_name: str,
    exam_name: str,
    report_date: str | None = None,
) -> dict:
    """构建报告文字上下文；仅使用页面已经得出的统计和名单。"""
    class_name = str(selected_class).strip() or "全部班级"
    subject_name = str(score_col).strip() or "未填写"
    school_display = str(school_name).strip() or "未填写"
    exam_display = str(exam_name).strip() or "成绩分析"

    return {
        "school_name": school_display,
        "exam_name": exam_display,
        "class_name": class_name,
        "subject_name": subject_name,
        "report_date": report_date or date.today().isoformat(),
        "student_count": str(analysis_result["student_count"]),
        "average_score": f"{float(analysis_result['average_score']):.2f}",
        "highest_score": _format_score(analysis_result["highest_score"]),
        "lowest_score": _format_score(analysis_result["lowest_score"]),
        "pass_rate": _format_rate(analysis_result["pass_rate"]),
        "excellent_rate": _format_rate(analysis_result["excellent_rate"]),
        "summary_text": _build_summary_text(analysis_result, class_name, subject_name, full_score),
        "distribution_analysis": _build_distribution_analysis(distribution),
        "level_analysis": _build_level_analysis(distribution),
        "excellent_students": _format_student_list(excellent_df, "暂无优秀学生"),
        "struggling_students": _format_student_list(fail_df, "暂无待提升学生"),
        "teaching_suggestions": _build_teaching_suggestions(analysis_result, distribution, subject_name),
    }


def safe_report_filename(
    *, school_name: str, class_name: str, subject_name: str, exam_name: str
) -> str:
    """构造 Windows 与 Streamlit 下载都可使用的报告文件名。"""

    def clean_part(value: str, fallback: str) -> str:
        cleaned = _ILLEGAL_FILENAME_CHARS.sub("", str(value)).strip().strip(".")
        return cleaned or fallback

    parts = [
        clean_part(school_name, "未填写"),
        clean_part(class_name, "全部班级"),
        clean_part(subject_name, "未填写"),
        clean_part(exam_name, "成绩分析"),
        "成绩分析报告",
    ]
    return "_".join(parts) + ".docx"


def _export_chart_images(distribution_figure: object, level_figure: object, output_dir: Path) -> tuple[Path, Path]:
    """将页面已有的图表导出为临时 PNG，不在项目目录写入文件。"""
    distribution_path = output_dir / "distribution.png"
    level_path = output_dir / "level.png"
    try:
        distribution_figure.write_image(distribution_path, format="png", width=1000, height=760, scale=2)
        level_figure.write_image(level_path, format="png", width=1000, height=800, scale=2)
    except Exception as exc:
        LOGGER.exception("Word report chart export failed: %s", exc)
        raise ReportGenerationError("Word 报告图表生成失败，请稍后重试或联系管理员。") from exc
    return distribution_path, level_path


def _ensure_no_unrendered_placeholders(report_bytes: bytes) -> None:
    with ZipFile(BytesIO(report_bytes)) as archive:
        xml_parts = [
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        ]
    if any(b"{{" in content or b"}}" in content for content in xml_parts):
        LOGGER.error("Word report rendering left template placeholders")
        raise ReportGenerationError("Word 报告模板存在未替换内容，请联系管理员。")


def build_score_report_bytes(
    *,
    analysis_result: dict,
    excellent_df: pd.DataFrame,
    fail_df: pd.DataFrame,
    distribution: pd.DataFrame,
    distribution_figure: object,
    level_figure: object,
    selected_class: str,
    score_col: str,
    full_score: float,
    school_name: str,
    exam_name: str,
    report_date: str | None = None,
    template_path: Path | None = None,
) -> bytes:
    """依次构建文字、导出图表、渲染并返回 Word 文件字节。"""
    template_file = Path(template_path) if template_path is not None else DEFAULT_TEMPLATE_PATH
    if not template_file.is_file():
        raise ReportGenerationError("Word 报告模板不存在，请联系管理员检查模板文件。")

    context = build_report_context(
        analysis_result=analysis_result,
        excellent_df=excellent_df,
        fail_df=fail_df,
        distribution=distribution,
        selected_class=selected_class,
        score_col=score_col,
        full_score=full_score,
        school_name=school_name,
        exam_name=exam_name,
        report_date=report_date,
    )

    with TemporaryDirectory(prefix="grade-report-") as temporary_dir:
        distribution_path, level_path = _export_chart_images(
            distribution_figure, level_figure, Path(temporary_dir)
        )
        try:
            template = DocxTemplate(template_file)
            # 模板正文宽度为 166 mm；保留安全边距并让两张图按原比例缩放。
            context["distribution_chart"] = InlineImage(template, str(distribution_path), width=Mm(158))
            context["level_chart"] = InlineImage(template, str(level_path), width=Mm(158))
            template.render(context, autoescape=True)
            output = BytesIO()
            template.save(output)
            report_bytes = output.getvalue()
            _ensure_no_unrendered_placeholders(report_bytes)
            return report_bytes
        except ReportGenerationError:
            raise
        except Exception as exc:
            LOGGER.exception("Word report rendering failed: %s", exc)
            raise ReportGenerationError("Word 报告生成失败，请稍后重试或联系管理员。") from exc
